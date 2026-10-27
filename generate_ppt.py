from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Colors (Clean Professional Palette)
    c_dark_bg = RGBColor(15, 23, 42)      # Slate 900
    c_light_bg = RGBColor(248, 250, 252)  # Slate 50
    c_white = RGBColor(255, 255, 255)
    c_navy = RGBColor(30, 41, 59)         # Slate 800
    c_primary = RGBColor(37, 99, 235)     # Blue 600
    c_accent = RGBColor(14, 165, 233)     # Sky 500
    c_text_dark = RGBColor(30, 41, 59)
    c_text_light = RGBColor(241, 245, 249)
    c_card_bg = RGBColor(255, 255, 255)
    c_card_border = RGBColor(226, 232, 240)

    def add_header(slide, title_text, category_text="ADVANCED MLOPS TOPIC"):
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(1.1))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        p0 = tf.paragraphs[0]
        p0.text = category_text.upper()
        p0.font.size = Pt(11)
        p0.font.bold = True
        p0.font.color.rgb = c_primary
        p0.space_after = Pt(4)

        p1 = tf.add_paragraph()
        p1.text = title_text
        p1.font.size = Pt(28)
        p1.font.bold = True
        p1.font.color.rgb = c_navy

    def set_slide_background(slide, color):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = color

    # ==========================================
    # SLIDE 1: Title Slide (Dark Theme)
    # ==========================================
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide1, c_dark_bg)

    accent_bar = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.2), Inches(2.2), Inches(0.15), Inches(2.8))
    accent_bar.fill.solid()
    accent_bar.fill.fore_color.rgb = c_accent
    accent_bar.line.fill.background()

    tb1 = slide1.shapes.add_textbox(Inches(1.6), Inches(2.1), Inches(10.5), Inches(3.2))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    
    p = tf1.paragraphs[0]
    p.text = "ADVANCED MLOPS PRESENTATION"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = c_accent
    p.space_after = Pt(10)

    p2 = tf1.add_paragraph()
    p2.text = "Automated Retraining in MLOps"
    p2.font.size = Pt(36)
    p2.font.bold = True
    p2.font.color.rgb = c_text_light
    p2.space_after = Pt(15)

    p3 = tf1.add_paragraph()
    p3.text = "How to Keep Machine Learning Models Accurate Automatically Over Time"
    p3.font.size = Pt(16)
    p3.font.color.rgb = RGBColor(148, 163, 184)

    tb_foot = slide1.shapes.add_textbox(Inches(1.6), Inches(6.2), Inches(10), Inches(0.5))
    tf_foot = tb_foot.text_frame
    pf = tf_foot.paragraphs[0]
    pf.text = "Presented for Advanced MLOps & Production AI Engineering"
    pf.font.size = Pt(12)
    pf.font.color.rgb = RGBColor(148, 163, 184)


    # ==========================================
    # SLIDE 2: Executive Summary & Context
    # ==========================================
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide2, c_light_bg)
    add_header(slide2, "Why Automated Retraining?", "Introduction")

    card_width = Inches(3.64)
    card_height = Inches(4.8)
    card_y = Inches(1.8)
    card_spacing = Inches(0.4)
    left_margin = Inches(0.8)

    cards_data_2 = [
        ("The Model Decay Problem", "Unlike traditional code that stays the same, machine learning models get outdated quickly because the real world changes (e.g., customer behavior, market trends, seasons)."),
        ("The Manual Bottleneck", "Traditionally, data scientists manually notice a drop in accuracy, pull new data, retrain in notebooks, and push code. This takes weeks and is prone to human error."),
        ("Automated Retraining", "Automating the feedback loop so that models detect performance drops and retrain themselves without human intervention ensures continuous accuracy and reliability.")
    ]

    for i, (ctitle, cdesc) in enumerate(cards_data_2):
        cx = left_margin + i * (card_width + card_spacing)
        shape = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, card_y, card_width, card_height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = c_card_bg
        shape.line.color.rgb = c_card_border
        shape.line.width = Pt(1)

        tf = shape.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.3)
        tf.margin_top = Inches(0.3)

        p = tf.paragraphs[0]
        p.text = ctitle
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.color.rgb = c_primary
        p.space_after = Pt(12)

        p_desc = tf.add_paragraph()
        p_desc.text = cdesc
        p_desc.font.size = Pt(14)
        p_desc.font.color.rgb = c_text_dark
        p_desc.space_before = Pt(4)


    # ==========================================
    # SLIDE 3: Concept & Triggers
    # ==========================================
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide3, c_light_bg)
    add_header(slide3, "Core Concept: When Should a Model Retrain?", "Conceptual Framework")

    steps = [
        ("1. Scheduled Retraining", "Retraining on a fixed calendar basis (e.g., every Sunday at midnight). Best for highly seasonal data like weekly sales or monthly forecasts."),
        ("2. Performance-Based Triggers", "Triggered automatically when production accuracy or F1-score drops below a defined threshold due to incoming ground truth labels."),
        ("3. Data Drift Triggers", "Triggered when incoming input data changes significantly compared to training data, even before model accuracy drops."),
        ("4. Event-Driven Triggers", "Triggered by major external business events (e.g., sudden tax law changes, new product launches, or market disruptions).")
    ]

    for i, (stitle, sdesc) in enumerate(steps):
        row = i // 2
        col = i % 2
        cx = Inches(0.8) if col == 0 else Inches(6.9)
        cy = Inches(1.8) if row == 0 else Inches(4.5)
        c_w = Inches(5.6)
        c_h = Inches(2.4)

        shape = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, cy, c_w, c_h)
        shape.fill.solid()
        shape.fill.fore_color.rgb = c_card_bg
        shape.line.color.rgb = c_card_border

        tf = shape.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.3)
        tf.margin_top = Inches(0.25)

        p = tf.paragraphs[0]
        p.text = stitle
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = c_primary
        p.space_after = Pt(8)

        p2 = tf.add_paragraph()
        p2.text = sdesc
        p2.font.size = Pt(13)
        p2.font.color.rgb = c_text_dark


    # ==========================================
    # SLIDE 4: Architecture & Workflow
    # ==========================================
    slide4 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide4, c_light_bg)
    add_header(slide4, "Architecture: The Retraining Pipeline", "System Architecture")

    arch_blocks = [
        ("1. Monitor & Detect", "Production monitoring tools continuously evaluate incoming predictions and data distributions, firing an alert when drift or decay occurs."),
        ("2. Pipeline Trigger", "An orchestrator (like Apache Airflow or Kubeflow) catches the alert and automatically kicks off the automated retraining pipeline."),
        ("3. Data Prep & Training", "The pipeline pulls fresh data from data lakes, runs feature engineering, trains the model, and validates performance against a test set."),
        ("4. Automated Deployment", "If the new model outperforms the current production model, CI/CD pipelines safely promote it to production via canary deployment.")
    ]

    for i, (atitle, adesc) in enumerate(arch_blocks):
        cy = Inches(1.8 + i * 1.3)
        shape = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), cy, Inches(11.7), Inches(1.1))
        shape.fill.solid()
        shape.fill.fore_color.rgb = c_card_bg
        shape.line.color.rgb = c_card_border

        tf = shape.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.4)
        tf.margin_right = Inches(0.4)
        tf.margin_top = Inches(0.15)

        p = tf.paragraphs[0]
        p.text = atitle
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = c_primary
        p.space_after = Pt(3)

        p2 = tf.add_paragraph()
        p2.text = adesc
        p2.font.size = Pt(13)
        p2.font.color.rgb = c_text_dark


    # ==========================================
    # SLIDE 5: Industry Toolchain
    # ==========================================
    slide5 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide5, c_light_bg)
    add_header(slide5, "The Automated Retraining Toolchain", "Tools & Ecosystem")

    tools_categories = [
        ("Orchestration & Pipelines", "Apache Airflow, Kubeflow Pipelines, Prefect, Argo Workflows"),
        ("Model Registry & Tracking", "MLflow, Weights & Biases (W&B), DVC (Data Version Control)"),
        ("Monitoring & Drift Detection", "Evidently AI, Arize AI, Deepchecks, WhyLabs"),
        ("Feature Stores", "Feast, Hopsworks, AWS Feature Store for consistent training/serving data"),
        ("CI/CD & Deployment", "GitHub Actions, ArgoCD, BentoML, Seldon Core")
    ]

    for i, (tcat, titems) in enumerate(tools_categories):
        cy = Inches(1.8 + i * 1.05)
        
        s_left = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), cy, Inches(3.6), Inches(0.9))
        s_left.fill.solid()
        s_left.fill.fore_color.rgb = c_navy
        s_left.line.fill.background()
        tf_l = s_left.text_frame
        tf_l.word_wrap = True
        tf_l.margin_left = tf_l.margin_right = Inches(0.2)
        tf_l.margin_top = Inches(0.25)
        p_l = tf_l.paragraphs[0]
        p_l.text = tcat
        p_l.font.size = Pt(14)
        p_l.font.bold = True
        p_l.font.color.rgb = c_white

        s_right = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.6), cy, Inches(7.9), Inches(0.9))
        s_right.fill.solid()
        s_right.fill.fore_color.rgb = c_card_bg
        s_right.line.color.rgb = c_card_border
        tf_r = s_right.text_frame
        tf_r.word_wrap = True
        tf_r.margin_left = tf_r.margin_right = Inches(0.3)
        tf_r.margin_top = Inches(0.25)
        p_r = tf_r.paragraphs[0]
        p_r.text = titems
        p_r.font.size = Pt(14)
        p_r.font.color.rgb = c_text_dark


    # ==========================================
    # SLIDE 6: Real-World Applications
    # ==========================================
    slide6 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide6, c_light_bg)
    add_header(slide6, "Real-World Applications", "Use Cases")

    apps = [
        ("E-Commerce Demand Forecasting", "Automatically retraining inventory prediction models weekly to adapt to sudden shopping trends, holidays, and flash sales."),
        ("Online Ad Click-Through Rate (CTR)", "Daily retraining of ad ranking models to incorporate fresh user engagement data and trending search queries."),
        ("Energy Grid Load Prediction", "Retraining weather-dependent power consumption models to handle heatwaves or cold snaps accurately."),
        ("Online Fraud Detection", "Updating fraud detection models rapidly as cybercriminals invent new scam patterns and transaction behaviors.")
    ]

    for i, (atitle, adesc) in enumerate(apps):
        row = i // 2
        col = i % 2
        cx = Inches(0.8) if col == 0 else Inches(6.9)
        cy = Inches(1.8) if row == 0 else Inches(4.5)
        c_w = Inches(5.6)
        c_h = Inches(2.4)

        shape = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, cy, c_w, c_h)
        shape.fill.solid()
        shape.fill.fore_color.rgb = c_card_bg
        shape.line.color.rgb = c_card_border

        tf = shape.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.3)
        tf.margin_top = Inches(0.25)

        p = tf.paragraphs[0]
        p.text = atitle
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = c_primary
        p.space_after = Pt(8)

        p2 = tf.add_paragraph()
        p2.text = adesc
        p2.font.size = Pt(13)
        p2.font.color.rgb = c_text_dark


    # ==========================================
    # SLIDE 7: Key Challenges
    # ==========================================
    slide7 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide7, c_light_bg)
    add_header(slide7, "Challenges & Risks of Automated Retraining", "Risk Management")

    challenges = [
        ("Feedback Loops & Bad Data", "If a newly retrained model learns from corrupted or biased data, automated pipelines can propagate bad models directly into production without human review."),
        ("Catastrophic Forgetting", "Retraining on recent data only can cause the model to 'forget' rare historical patterns or edge cases that happen infrequently."),
        ("Compute Costs & Over-Retraining", "Retraining large models too frequently is expensive and unnecessary if data patterns haven't genuinely changed."),
        ("Silent Failures", "If training data pipelines break upstream, automated retraining jobs might complete successfully with garbage inputs, producing broken models.")
    ]

    for i, (ctitle, cdesc) in enumerate(challenges):
        row = i // 2
        col = i % 2
        cx = Inches(0.8) if col == 0 else Inches(6.9)
        cy = Inches(1.8) if row == 0 else Inches(4.5)
        c_w = Inches(5.6)
        c_h = Inches(2.4)

        shape = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, cy, c_w, c_h)
        shape.fill.solid()
        shape.fill.fore_color.rgb = c_card_bg
        shape.line.color.rgb = c_card_border

        tf = shape.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.3)
        tf.margin_top = Inches(0.25)

        p = tf.paragraphs[0]
        p.text = ctitle
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = RGBColor(220, 38, 38)
        p.space_after = Pt(8)

        p2 = tf.add_paragraph()
        p2.text = cdesc
        p2.font.size = Pt(13)
        p2.font.color.rgb = c_text_dark


    # ==========================================
    # SLIDE 8: Conclusion & Future Outlook
    # ==========================================
    slide8 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide8, c_dark_bg)

    tb8 = slide8.shapes.add_textbox(Inches(1.0), Inches(1.2), Inches(11.3), Inches(5.2))
    tf8 = tb8.text_frame
    tf8.word_wrap = True

    p = tf8.paragraphs[0]
    p.text = "CONCLUSION & FUTURE OUTLOOK"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = c_accent
    p.space_after = Pt(10)

    p2 = tf8.add_paragraph()
    p2.text = "Building Self-Healing AI Systems"
    p2.font.size = Pt(32)
    p2.font.bold = True
    p2.font.color.rgb = c_text_light
    p2.space_after = Pt(20)

    bullet_points = [
        "From Static to Adaptive: Automated retraining ensures models evolve alongside real-world changes without manual overhead.",
        "Human-in-the-Loop Safety: Combining automated pipelines with validation gates ensures speed without sacrificing quality and safety.",
        "The Future: Fully autonomous MLOps loops where models continuously monitor, learn, test, and deploy safely."
    ]

    for bp in bullet_points:
        p_bp = tf8.add_paragraph()
        p_bp.text = "• " + bp
        p_bp.font.size = Pt(16)
        p_bp.font.color.rgb = RGBColor(203, 213, 225)
        p_bp.space_after = Pt(12)

    prs.save("Advanced_MLOps_Automated_Retraining_Presentation.pptx")
    print("Successfully generated Advanced_MLOps_Automated_Retraining_Presentation.pptx")

if __name__ == "__main__":
    create_presentation()
